"""XLSX renderer for the content-gap report: one sheet per report section, plus
full-detail sheets (every heading, section text, FAQ, link and image).

render(ctx, path) -> path, or None if openpyxl is missing (caller falls back to CSV).
Styling is deliberately plain: one navy header row, frozen panes, filters,
no gridlines, and a light fill on your page's row or column only.
"""
from datetime import date

from report_core import (DEPTH_LABEL, FIELD_CATEGORY, PRIORITY_LABEL, _norm, clean_text, internal_link_count, keyword_checks,
                         map_get, onpage_internal_links, serp_matrix, sim_verdict, similarity)

NAVY, INK, MUTED = "1E3A5F", "0F172A", "64748B"
FIELD_WORD = {"good": "Good", "ni": "Needs improvement", "poor": "Poor"}


def cell_value(v):
    """Excel takes scalars only: a person record becomes its name, lists are joined."""
    if isinstance(v, dict):
        return str(v.get("name") or v.get("title") or ", ".join(f"{k}: {x}" for k, x in v.items()))
    if isinstance(v, (list, tuple, set)):
        return "; ".join(str(cell_value(x)) for x in v)
    return v


def render(ctx, path):
    try:
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, Reference
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        return None

    c = ctx
    order, your, pages, topic = c["order"], c["your"], c["pages"], c["topic"]
    thin = Side(style="thin", color="E2E8F0")
    F = lambda **k: Font(name="Segoe UI", **k)
    hdr_font, hdr_fill = F(bold=True, color="FFFFFF", size=10), PatternFill("solid", fgColor=NAVY)
    body_font, you_fill = F(size=10, color="334155"), PatternFill("solid", fgColor="F1F5F9")
    yn = lambda v: "Yes" if v else "No"
    wb = Workbook()
    made = []

    def sheet(name, heading, sub, headers, rows, widths, you_col=None, you_row_col=None, wrap=()):
        ws = wb.create_sheet(name)
        made.append(name)
        ws.sheet_view.showGridLines = False
        ws["A1"], ws["A2"] = heading, sub
        ws["A1"].font, ws["A2"].font = F(size=14, bold=True, color=INK), F(size=9, color=MUTED)
        for i, h in enumerate(headers, 1):
            cell = ws.cell(row=4, column=i, value=h)
            cell.font, cell.fill, cell.alignment = hdr_font, hdr_fill, Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[4].height = 30
        for r, rv in enumerate(rows, 5):
            is_you = you_row_col is not None and rv[you_row_col] == your
            for ci, v in enumerate(rv, 1):
                v = cell_value(v)
                if isinstance(v, str) and len(v) > 32000:
                    v = v[:32000] + " …"
                cell = ws.cell(row=r, column=ci, value=v)
                cell.font, cell.border = body_font, Border(bottom=thin)
                cell.alignment = Alignment(vertical="top", wrap_text=(ci - 1) in wrap)
                if is_you or (you_col is not None and ci - 1 == you_col):
                    cell.fill = you_fill
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "B5"
        ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{4 + max(1, len(rows))}"
        return ws

    # ---- Summary ------------------------------------------------------------
    kpi, qual, psi_by, auth_by = c["kpis"], c["quality"], c["psi_by"], c["auth_by"]
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    ws["A1"] = f"SEO Competitive Audit: {topic[:1].upper() + topic[1:]}"
    ws["A1"].font = F(size=16, bold=True, color=INK)
    ws["A2"] = f'{pages[your].get("url")} · {date.today().strftime("%d %b %Y")}'
    ws["A2"].font = F(size=9, color=MUTED)
    m = ((psi_by.get(your) or {}).get("mobile") or {})
    fld = m.get("field") or {}
    a = auth_by.get(your) or {}
    rk = c["ranking"].get("our_ranking") or {}
    kp = [("Content coverage", f'{kpi.get("coverage_pct", "—")}%', f'{kpi.get("missing_count", 0)} missing · {kpi.get("thin_count", 0)} thin'),
          ("Content quality", (map_get(qual, your, {}) or {}).get("quality_score"), "/ 100"),
          ("FAQs on your page", len(pages[your].get("faqs") or []), f'{kpi.get("faq_gap_count", 0)} competitor questions unanswered'),
          ("Core Web Vitals (mobile, real users)",
           FIELD_WORD.get(FIELD_CATEGORY.get(fld.get("overall_category") or "", ""), "No data") if m and not m.get("error") else "Not available", ""),
          ("Mobile performance (lab)", (m.get("scores") or {}).get("performance") if m and not m.get("error") else "Not available", "/ 100"),
          ("URL Rating (Ahrefs)", a.get("url_rating") if a else "Not available", f'{a.get("refdomains")} referring domains' if a else ""),
          ("Organic keywords (Ahrefs)", a.get("org_keywords") if a else "Not available", f'{a.get("org_keywords_top3")} in top 3' if a else ""),
          ("Google ranking (web search)", f'#{rk.get("best_position")}' if rk.get("found") else "Not ranked", f'{len(c["ranking"].get("queries_used") or [])} queries checked')]
    for i, h in enumerate(["Metric", "Value", "Context"], 1):
        cell = ws.cell(row=4, column=i, value=h)
        cell.font, cell.fill = hdr_font, hdr_fill
    for r, row in enumerate(kp, 5):
        for ci, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=ci, value=v)
            cell.font, cell.border = F(size=10, bold=ci == 2, color=INK), Border(bottom=thin)
    for col, w in zip("ABC", (34, 20, 46)):
        ws.column_dimensions[col].width = w

    # ---- Page speed -----------------------------------------------------------
    if psi_by:
        rows = []
        for b in order:
            for dev in ("mobile", "desktop"):
                r = (psi_by.get(b) or {}).get(dev) or {}
                if not r or r.get("error"):
                    rows.append([b, dev.title(), "Not measured", (r or {}).get("error", "")[:200]] + [None] * 18)
                    continue
                f, lab, sc = r.get("field") or {}, r.get("lab") or {}, r.get("scores") or {}
                rows.append([b, dev.title(), FIELD_WORD.get(FIELD_CATEGORY.get(f.get("overall_category") or "", ""), "No data"),
                             ("Site-wide" if f.get("origin_fallback") else "This page") if f.get("has_data") else "",
                             f.get("lcp_ms"), f.get("inp_ms"), f.get("cls"), f.get("fcp_ms"), f.get("ttfb_ms"),
                             sc.get("performance"), lab.get("lcp_ms"), lab.get("cls"), lab.get("tbt_ms"), lab.get("fcp_ms"), lab.get("si_ms"),
                             lab.get("ttfb_ms"), round((lab.get("bytes") or 0) / 1048576, 2), sc.get("accessibility"), sc.get("best-practices"), sc.get("seo"),
                             "; ".join(fx["title"] for fx in (r.get("fixes") or [])[:5]), r.get("fetched_at")])
        sheet("Page Speed", "Page speed & Core Web Vitals", "Google PageSpeed Insights · real users = CrUX p75, 28 days · times in ms",
              ["Page", "Device", "CWV assessment (real users)", "Real-user data scope", "LCP (real)", "INP (real)", "CLS (real)", "FCP (real)", "TTFB (real)",
               "Performance (lab)", "LCP (lab)", "CLS (lab)", "TBT (lab)", "FCP (lab)", "Speed Index (lab)", "Server response (lab)", "Page weight (MB)",
               "Accessibility", "Best practices", "SEO", "Top Lighthouse findings", "Measured at"],
              rows, [20, 9, 18, 14] + [10] * 16 + [70, 22], you_row_col=0, wrap=(20,))
        sheet("Speed Fixes", "Lighthouse findings per page", "Audits scoring below 90 with an estimated saving or affected metric",
              ["Page", "Device", "Finding", "Metrics affected", "Est. saving (ms)", "Detail"],
              [[b, dev.title(), fx.get("title"), ", ".join(fx.get("metrics") or []), fx.get("savings_ms") or None, fx.get("display")]
               for b in order for dev in ("mobile", "desktop") for fx in (((psi_by.get(b) or {}).get(dev) or {}).get("fixes") or [])],
              [20, 9, 50, 18, 14, 40], you_row_col=0, wrap=(2,))
    else:
        sheet("Page Speed", "Page speed & Core Web Vitals", "Not available yet: PageSpeed Insights was not run for this report",
              ["Page", "Status"], [[b, "Not available"] for b in order], [24, 20], you_row_col=0)

    # ---- Technical --------------------------------------------------------------
    P = lambda b: pages[b]
    tech = [("Title length", lambda b: len(P(b).get("title") or "")), ("Meta description length", lambda b: len(P(b).get("meta_description") or "")),
            ("H1 count", lambda b: c["hcounts"][b][1]), ("H2 count", lambda b: c["hcounts"][b][2]), ("H3 count", lambda b: c["hcounts"][b][3]),
            ("H4 count", lambda b: c["hcounts"][b][4]), ("Visible words", lambda b: P(b).get("word_count_total")), ("Tables", lambda b: P(b).get("tables_count")),
            ("Images", lambda b: P(b).get("image_count")),
            ("Images missing alt", lambda b: max(0, int(P(b).get("image_count") or 0) - int(P(b).get("image_alt_count") or 0))),
            ("Structured-data types", lambda b: ", ".join(P(b).get("schema_types") or [])),
            ("FAQPage schema", lambda b: yn("FAQPage" in (P(b).get("schema_types") or []))),
            ("Named author", lambda b: P(b).get("author") or "No"), ("Named reviewer", lambda b: P(b).get("reviewer") or "No"),
            ("Published", lambda b: P(b).get("published_date") or "Not shown"), ("Last updated", lambda b: P(b).get("modified_date") or "Not shown"),
            ("Internal links (article body)", lambda b: internal_link_count(P(b), qual)), ("External links", lambda b: P(b).get("external_link_count"))]
    sheet("Technical SEO", "Technical & on-page SEO", "From raw HTML", ["Check"] + order, [[k] + [fn(b) for b in order] for k, fn in tech],
          [28] + [22] * len(order), you_col=1, wrap=tuple(range(1, len(order) + 1)))
    if qual:
        sheet("Content Quality", "Content quality signals", "From the gap analysis",
              ["Page", "Quality score", "Words", "H2", "FAQs", "Tables", "Internal links", "Schema", "E-E-A-T", "Freshness", "Note"],
              [[b, q.get("quality_score"), q.get("word_count"), q.get("h2"), q.get("faqs"), q.get("tables"), q.get("internal_links"),
                ", ".join(q.get("schema") or []), yn(q.get("eeat")), q.get("freshness"), q.get("note", "")] for b in order for q in [map_get(qual, b, {}) or {}]],
              [20, 12, 9, 6, 6, 7, 12, 40, 8, 40, 60], you_row_col=0, wrap=(7, 9, 10))

    # ---- Keywords ---------------------------------------------------------------
    kc = {b: keyword_checks(pages[b], topic) for b in order}
    sheet("Keywords", f"Keyword placement: “{topic}”", "Yes = every word of the phrase appears", ["Placement"] + order,
          [["Title tag"] + [yn(kc[b]["title"]) for b in order], ["Meta description"] + [yn(kc[b]["meta"]) for b in order],
           ["H1"] + [yn(kc[b]["h1"]) for b in order], ["URL"] + [yn(kc[b]["url"]) for b in order],
           ["First 100 words"] + [yn(kc[b]["first100"]) for b in order],
           ["H2/H3 headings containing it"] + [f'{kc[b]["subheads"]} of {kc[b]["subheads_total"]}' for b in order],
           ["Exact phrase in body"] + [kc[b]["exact_body"] for b in order], ["Exact phrase per 1,000 words"] + [kc[b]["per_1000"] for b in order]],
          [30] + [18] * len(order), you_col=1)
    sheet("Titles & Meta", "Title, meta description and H1", "Verbatim", ["Page", "Title", "Title chars", "Meta description", "Meta chars", "H1", "URL"],
          [[b, P(b).get("title"), len(P(b).get("title") or ""), P(b).get("meta_description"), len(P(b).get("meta_description") or ""), P(b).get("h1"), P(b).get("url")] for b in order],
          [20, 60, 11, 70, 11, 34, 60], you_row_col=0, wrap=(1, 3))
    if auth_by:
        sheet("Ranking Keywords", "Keywords each page ranks for (Ahrefs)", f'Google {(c["auth"] or {}).get("country", "")} · top keywords by traffic',
              ["Page", "Keyword", "Position", "Volume", "Traffic"],
              [[b, k.get("keyword"), k.get("position"), k.get("volume"), k.get("traffic")] for b in order for k in (auth_by.get(b) or {}).get("top_keywords") or []],
              [20, 50, 10, 10, 10], you_row_col=0)
    qs, best = serp_matrix(c["ranking"])
    if qs:
        sheet("Search Ranking", "Search ranking by query", "Web search · blank = not found", ["Query"] + order,
              [[q] + [best.get((q, b)) for b in order] for q in qs], [46] + [18] * len(order), you_col=1)
        sheet("Top Results", "Top results per query", "", ["Query", "Rank", "Domain", "Title", "URL", "Same page type"],
              [[r.get("query"), r.get("rank"), r.get("domain") or r.get("brand"), r.get("title"), r.get("url"), yn(r.get("same_type"))]
               for r in c["ranking"].get("ranking_pages") or []], [40, 7, 24, 60, 70, 12], wrap=(3,))

    # ---- Visibility -------------------------------------------------------------
    if auth_by:
        sheet("Authority", "Search visibility & authority (Ahrefs)", f'Exact URL · Domain Rating for the whole site · {(c["auth"] or {}).get("fetched_at", "")[:10]}',
              ["Page", "URL Rating", "Domain Rating", "Organic traffic", "Keywords (top 100)", "Keywords (top 3)", "Referring domains", "Backlinks", "URL"],
              [[b] + [(auth_by.get(b) or {}).get(k) for k in ("url_rating", "domain_rating", "org_traffic", "org_keywords", "org_keywords_top3", "refdomains", "backlinks")]
               + [P(b).get("url")] for b in order], [20, 11, 13, 15, 17, 15, 17, 12, 60], you_row_col=0)
    else:
        sheet("Authority", "Search visibility & authority", "Not available yet: Ahrefs was not connected for this report",
              ["Page", "Status"], [[b, "Not available"] for b in order], [24, 20], you_row_col=0)
    if c["gaps"].get("ranking_assessment"):
        sheet("Ranking Assessment", "Ranking assessment", "", ["Page", "Google", "AI search"],
              [[b, a2.get("google"), a2.get("ai_search")] for b in order for a2 in [map_get(c["gaps"]["ranking_assessment"], b, {}) or {}]],
              [20, 90, 70], you_row_col=0, wrap=(1, 2))
    sheet("External Mentions", "External sources and brands cited", "", ["Page", "Mention"],
          [[b, x] for b in order for x in map_get(c["gaps"].get("external_brands") or {}, b, []) or []], [20, 90], you_row_col=0)

    # ---- Headings ---------------------------------------------------------------
    hc = c["hcounts"]
    sheet("Heading Levels", "Heading counts H1–H4", "", ["Level"] + order, [[f"H{l}"] + [hc[b][l] for b in order] for l in (1, 2, 3, 4)],
          [12] + [18] * len(order), you_col=1)
    sheet("Heading Hierarchy", "Heading hierarchy H1 › H2 › H3 › H4", "Yes = heading present on that page", ["Level", "Heading"] + order + ["Pages using it"],
          [[f"H{r['level']}", "    " * (r["level"] - 1) + r["text"]] + [yn(r["present"][b]) for b in order] + [sum(r["present"].values())] for r in c["tree"]],
          [8, 70] + [16] * len(order) + [13], you_col=2)
    nih = [[b, r["text"], "<%s>" % r["tag"], r["parent_heading"], r["suggested"]] for b in order for r in c["not_in_headings"][b]]
    if nih:
        sheet("Heading Markup Gaps", "Styled as headings but not H-tags", "", ["Page", "Text", "Current tag", "Under heading", "Should be"], nih,
              [20, 60, 12, 40, 10], you_row_col=0, wrap=(1,))
    sheet("Page Structure", "Every heading and its content, in page order", "", ["Page", "#", "Level", "Heading", "Words", "Content"],
          [[b, i, f'H{s.get("level")}', s.get("heading"), s.get("word_count"), s.get("text")]
           for b in order for i, s in enumerate(pages[b].get("sections") or [], 1)], [20, 5, 7, 44, 8, 110], you_row_col=0, wrap=(3, 5))

    # ---- Topics -----------------------------------------------------------------
    sheet("Topic Coverage", "Topic coverage: heading level used and depth", "— = not covered · Body = no dedicated heading", ["Topic", "Intent"] + order,
          [[t["name"], t["intent"]] + [(f'{t["brands"][b]["tag"]} · {DEPTH_LABEL.get(int(t["brands"][b]["info"].get("depth", 0) or 0), "")}'
                                        if t["brands"][b]["present"] else "—") for b in order] for t in c["topics"]],
          [50, 14] + [18] * len(order), you_col=2, wrap=(0,))
    sheet("Unique Coverage", "Topics only one page covers", "", ["Page", "Topic"], [[b, t] for b in order for t in c["unique"].get(b) or []], [20, 70], you_row_col=0)
    rows = []
    for t in c["topics"]:
        yt = (t["brands"][your]["sec"] or {}).get("text") or ""
        for b in order:
            v = t["brands"][b]
            if not v["present"]:
                rows.append([t["name"], b, "—", "", "", "", "Not covered", "", "", ""])
                continue
            info, sec = v["info"], v["sec"] or {}
            txt = sec.get("text") or ""
            pct = similarity(yt, txt) if (b != your and yt and txt) else None
            rows.append([t["name"], b, v["tag"], sec.get("heading") or "", DEPTH_LABEL.get(int(info.get("depth", 0) or 0), ""), info.get("word_count"),
                         info.get("snippet") or "", pct, sim_verdict(pct) if pct is not None else "", txt])
    sheet("Topic Content", "Exact content by topic", "One row per topic × page",
          ["Topic", "Page", "Heading level", "Heading", "Depth", "Words", "Summary", "Similarity to yours %", "Similarity", "Section text"],
          rows, [34, 18, 9, 36, 10, 8, 50, 12, 16, 100], you_row_col=1, wrap=(0, 3, 6, 9))
    sheet("Gaps", "Prioritised gaps", c["kpis"].get("priority_method", ""), ["Priority", "Type", "Gap", "Detail", "Recommendation", "Best example", "Topic"],
          [[PRIORITY_LABEL.get(int(g.get("priority") or 0)), (g.get("type") or "").title(), g.get("title"), g.get("detail"), g.get("recommendation"),
            g.get("exemplar_brand"), g.get("cluster")] for g in c["gap_list"]], [9, 10, 50, 80, 70, 18, 24], wrap=(2, 3, 4))

    # ---- FAQs -------------------------------------------------------------------
    fg = c["gaps"].get("faq_gaps") or []
    if fg:
        sheet("FAQ Gaps", "Questions competitors answer and you don’t", "", ["Question"] + order + ["Competitors answering"],
              [[f.get("question")] + [yn(_norm(b) in {_norm(x) for x in f.get("answered_by") or []}) for b in order] + [len(f.get("answered_by") or [])] for f in fg],
              [70] + [16] * len(order) + [14], you_col=1, wrap=(0,))
    sheet("FAQ Coverage", "All questions by page", "Differently worded questions grouped", ["Question"] + order + ["Pages answering"],
          [[r["question"]] + [yn(r["present"][b]) for b in order] + [sum(r["present"].values())] for r in c["faq_rows"]],
          [70] + [16] * len(order) + [13], you_col=1, wrap=(0,))
    sheet("FAQs Verbatim", "Questions and answers as published", "", ["Page", "#", "Question", "Answer"],
          [[b, i, f.get("question"), clean_text(f.get("answer") or "")] for b in order for i, f in enumerate(pages[b].get("faqs") or [], 1)],
          [20, 5, 55, 100], you_row_col=0, wrap=(2, 3))

    # ---- Links & images ------------------------------------------------------------
    lg = c["gaps"].get("link_gaps") or []
    if lg:
        sheet("Link Gaps", "Internal-link gaps", "", ["Link target"] + order,
              [[l.get("topic_or_target")] + [yn(_norm(b) in {_norm(x) for x in l.get("present_in") or []}) for b in order] for l in lg],
              [60] + [16] * len(order), you_col=1, wrap=(0,))
    sheet("Links", "Internal and external links", "Internal = on-page editorial links (nav/footer excluded)", ["Page", "Type", "Anchor", "URL", "Section"],
          [[b, "Internal", l.get("anchor"), l.get("href"), l.get("section")] for b in order for l in onpage_internal_links(pages[b])]
          + [[b, "External", l.get("anchor"), l.get("href"), l.get("section")] for b in order for l in pages[b].get("external_links") or []],
          [20, 10, 40, 70, 40], you_row_col=0)
    sheet("Images", "Images", "", ["Page", "Alt text", "Source"], [[b, i.get("alt"), i.get("src")] for b in order for i in pages[b].get("images") or []],
          [20, 50, 90], you_row_col=0)
    top = [g for g in c["gap_list"] if g.get("recommendation")]
    top = [g for g in top if int(g.get("priority") or 0) == 3] or top[:10]
    sheet("Recommendations", "Recommendations", "Highest-priority gaps", ["#", "Priority", "Area", "Gap", "Recommendation", "Best example"],
          [[i, PRIORITY_LABEL.get(int(g.get("priority") or 0)), (g.get("type") or "").title(), g.get("title"), g.get("recommendation"), g.get("exemplar_brand")]
           for i, g in enumerate(top, 1)], [5, 9, 12, 50, 90, 18], wrap=(3, 4))

    # ---- Charts (monochrome) ------------------------------------------------------------
    cols = ["Page", "Words", "H2", "H3", "FAQs", "Internal links"]
    data = [[b, P(b).get("word_count_total"), hc[b][2], hc[b][3], len(P(b).get("faqs") or []), internal_link_count(P(b), qual)] for b in order]
    charts = [("Words", 2), ("FAQs", 5), ("H3 headings", 4)]
    if psi_by:
        cols.append("Mobile performance")
        for row, b in zip(data, order):
            row.append((((psi_by.get(b) or {}).get("mobile") or {}).get("scores") or {}).get("performance"))
        charts.append(("Mobile performance (lab)", len(cols)))
    if auth_by:
        cols.append("Referring domains")
        for row, b in zip(data, order):
            row.append((auth_by.get(b) or {}).get("refdomains"))
        charts.append(("Referring domains", len(cols)))
    ch = sheet("Charts", "Page comparison", "", cols, data, [22] + [12] * (len(cols) - 1), you_row_col=0)
    for i, (title, col) in enumerate(charts):
        bc = BarChart()
        bc.type, bc.title, bc.legend, bc.height, bc.width = "bar", title, None, 6.5, 13
        bc.add_data(Reference(ch, min_col=col, min_row=4, max_row=4 + len(order)), titles_from_data=True)
        bc.set_categories(Reference(ch, min_col=1, min_row=5, max_row=4 + len(order)))
        bc.series[0].graphicalProperties.solidFill = "94A3B8"
        bc.series[0].graphicalProperties.line.solidFill = "94A3B8"
        ch.add_chart(bc, f"{get_column_letter(len(cols) + 2)}{4 + i * 14}")

    psi, auth = c["psi"], c["auth"]
    sheet("Sources", "Method & data sources", "", ["Source", "What it provides", "Status"],
          [["Page crawl (raw HTML)", "Headings, content, FAQs, links, images, meta, schema", "Collected"],
           ["Gap analysis", "Topics, depth, gaps, priorities, quality", "Collected"],
           ["Web search", "Competitors, ranking check", "Collected" if c["ranking"] else "Not run"],
           ["Google PageSpeed Insights API v5", "Core Web Vitals (CrUX), Lighthouse", f'Collected {(psi.get("fetched_at") or "")[:10]}' if psi else "Not available yet"],
           ["Ahrefs API v3", "Authority, backlinks, organic keywords and traffic", f'Collected {(auth.get("fetched_at") or "")[:10]}' if auth else "Not available yet"],
           ["Semrush", "Keyword and traffic data", "Not included"]], [32, 50, 22])

    ws.cell(row=len(kp) + 6, column=1, value="Sheets").font = F(bold=True, color=INK)
    for i, s in enumerate(made):
        cell = ws.cell(row=len(kp) + 7 + i, column=1, value=s)
        cell.hyperlink, cell.font = f"#'{s}'!A1", F(size=10, color=NAVY, underline="single")
    wb.save(path)
    return path
